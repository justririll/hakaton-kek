"""Извлечение показателей Форм 1 и 2 из разнородных Excel-выгрузок.

Двадцать организаций присылают один и тот же шаблон, но с расхождениями:
лишние пробелы в имени листа, сдвинутые колонки, вставленные строки со
ссылками на подтверждающие документы, `.xls` вместо `.xlsx`, переименованная
колонка «Фактическое значение …» → «Значение показателя 2026 года».

Поэтому парсер не опирается на фиксированные координаты: он находит строку
заголовка, строит карту колонок по ключевым словам и извлекает строки по
номеру показателя, сверяя название с эталонными ключевыми словами схемы.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

from app.ingest.schema import BY_FORM_CODE, FORM1, FORM1_APP, FORM2, TITLE

log = logging.getLogger(__name__)

# Пробельные символы, которыми Excel любит разделять разряды числа.
_SPACES = "     \t"
_NUM_RE = re.compile(rf"^-?\d[\d{_SPACES}]*(?:[.,]\d+)?$")
_CODE_RE = re.compile(r"^(\d+(?:\.\d+)?)\s*[.)]?\s")
# «тыс. рублей», «тыс.рублей», «тысяч рублей» — часть организаций отчитывается так.
_THOUSAND_RE = re.compile(r"тыс[\s.]*(?:яч)?[\s.]*руб")


def _unit_hint(name: str) -> str:
    """Единица измерения денежной строки, объявленная в её названии."""
    return "thousand" if _THOUSAND_RE.search(name.lower()) else "unit"


def normalize(value: object) -> str:
    """Схлопывает пробелы и приводит к строке; None → пустая строка."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def to_number(value: object) -> float | None:
    """Парсит число из ячейки, терпя разрядные пробелы и запятую-разделитель."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize(value)
    if not text or not _NUM_RE.match(text):
        return None
    for space in _SPACES:
        text = text.replace(space, "")
    return float(text.replace(",", "."))


# --- чтение книги -------------------------------------------------------------

def read_workbook(path: Path) -> dict[str, list[list[object]]]:
    """Читает книгу (.xls/.xlsx) в словарь {имя листа: матрица значений}."""
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        return {
            sheet.name: [
                [sheet.cell_value(r, c) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)
            ]
            for sheet in book.sheets()
        }

    import openpyxl

    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        return {
            name: [[cell.value for cell in row] for row in book[name].iter_rows()]
            for name in book.sheetnames
        }
    finally:
        book.close()


def sheet_role(name: str) -> str | None:
    """Определяет роль листа по его названию."""
    upper = normalize(name).upper().replace("Ф.", "Ф")
    if "ТИТУЛ" in upper:
        return TITLE
    if "ПРИЛОЖ" in upper:
        return FORM1_APP
    if "ФОРМА 2" in upper or upper.endswith("Ф2"):
        return FORM2
    if "ФОРМА 1" in upper or upper.endswith("Ф1"):
        return FORM1
    return None


# --- карта колонок ------------------------------------------------------------

# Порядок важен: более специфичные шаблоны проверяются первыми.
_COLUMN_PATTERNS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("code", ("№ п/п", "№п/п")),
    ("name", ("наименование",)),
    ("plan_pct", ("плановое значение показателя, процент",)),
    ("baseline", ("значение показателя в 2025", "значение показателя 2025")),
    ("fact", (
        "фактическое значение показателя за отчетный период",
        "фактическое значение показателя",
        "значение показателя 2026",
    )),
    ("formula", ("формула",)),
    ("docs", ("документы",)),
    ("deviation", ("причина отклонения",)),
    ("value", ("значение показателя",)),  # Форма 2 — единственная колонка значения
)


def _find_header(rows: list[list[object]]) -> tuple[int, dict[str, int]] | None:
    """Ищет строку заголовка таблицы и строит карту {роль колонки: индекс}."""
    for index, row in enumerate(rows[:15]):
        cells = [normalize(cell).lower() for cell in row]
        if not any(c.startswith("наименование") for c in cells):
            continue
        mapping: dict[str, int] = {}
        for col, text in enumerate(cells):
            if not text:
                continue
            for role, patterns in _COLUMN_PATTERNS:
                if role in mapping:
                    continue
                if any(p in text for p in patterns):
                    mapping[role] = col
                    break
        if "name" in mapping:
            return index, mapping
    return None


def _row_code(row: list[object], columns: dict[str, int], name: str) -> str | None:
    """Номер показателя: из колонки «№ п/п» либо из префикса названия («3.1.»)."""
    raw = normalize(row[columns["code"]]) if "code" in columns and columns["code"] < len(row) else ""
    if raw:
        number = to_number(raw)
        if number is not None:
            return f"{number:g}"
        match = _CODE_RE.match(raw + " ")
        if match:
            return match.group(1)
    match = _CODE_RE.match(name + " ")
    return match.group(1) if match else None


def _pick_value(row: list[object], columns: dict[str, int], roles: tuple[str, ...]) -> float | None:
    """Первое непустое число из колонок, перечисленных в `roles`."""
    for role in roles:
        col = columns.get(role)
        if col is None or col >= len(row):
            continue
        number = to_number(row[col])
        if number is not None:
            return number
    return None


@dataclass
class ParsedValue:
    """Одно извлечённое значение показателя."""

    key: str
    code: str
    form: str
    fact: float | None
    baseline: float | None
    plan_pct: float | None
    raw_name: str
    name_matched: bool
    unit_hint: str = "unit"  # "thousand", если строка подписана «тыс. рублей»


def _parse_indicator_sheet(rows: list[list[object]], form: str) -> tuple[list[ParsedValue], list[str]]:
    """Разбирает лист одной формы. Возвращает значения и список предупреждений."""
    warnings: list[str] = []
    header = _find_header(rows)
    if header is None:
        return [], [f"{form}: не найдена строка заголовка"]

    header_row, columns = header
    seen: set[str] = set()
    values: list[ParsedValue] = []

    for row in rows[header_row + 1 :]:
        name_col = columns["name"]
        name = normalize(row[name_col]) if name_col < len(row) else ""
        code = _row_code(row, columns, name)
        if code is None or code in seen:
            continue
        indicator = BY_FORM_CODE.get((form, code))
        if indicator is None:
            continue
        # Строка-подпункт без названия в своей колонке — служебная, пропускаем.
        if not name:
            continue

        lowered = name.lower()
        matched = any(kw in lowered for kw in indicator.keywords) if indicator.keywords else True
        if not matched:
            warnings.append(
                f"{form}/{code}: название не совпало с эталоном → {name[:70]!r}"
            )

        seen.add(code)
        values.append(
            ParsedValue(
                unit_hint=_unit_hint(name) if indicator.monetary else "unit",
                key=indicator.key,
                code=code,
                form=form,
                fact=_pick_value(row, columns, ("fact", "value")),
                baseline=_pick_value(row, columns, ("baseline",)),
                plan_pct=_pick_value(row, columns, ("plan_pct",)),
                raw_name=name,
                name_matched=matched,
            )
        )

    expected = {code for (f, code) in BY_FORM_CODE if f == form}
    for code in sorted(expected - seen):
        warnings.append(f"{form}/{code}: строка показателя отсутствует")
    return values, warnings


# --- титульный лист -----------------------------------------------------------

_ORG_MARKERS = ("образовательная организация", "организация, на базе которой")
_CENTER_MARKERS = ("центр прототипирования", "творческий инкубатор")

# Служебные подписи, которые часть организаций оставляет слитно со значением
# (иногда без двоеточия): «Образовательная организация, … создан центр
# прототипирования Новосибирская …».
# Длинные подписи снимаются всегда: они не бывают частью названия.
_LONG_LABELS = (
    "образовательная организация, на базе которой создан центр прототипирования",
    "образовательная организация, на базе которой создан творческий инкубатор",
    "образовательная организация, на базе которой создан",
    "образовательная организация",
)
# Короткие — только вместе с двоеточием: иначе «Центр прототипирования и 3D
# технологий» потеряет половину собственного названия.
_SHORT_LABELS = ("центр прототипирования", "творческий инкубатор")


# Срок сдачи отчёта лежит в соседней колонке той же строки и при склейке
# прилипает к названию организации.
_TRAILING_NOISE = re.compile(
    r"\s*(?:не\s+позднее\s+)?\d{1,2}\s+(?:янв|фев|мар|апр|ма|июн|июл|авг|сен|окт|ноя|дек)\w*\s*$",
    re.IGNORECASE,
)


def _strip_label(text: str) -> str:
    """Убирает служебные подписи поля и хвост со сроком сдачи."""
    cleaned = _TRAILING_NOISE.sub("", text.strip())
    # Подпись встречается и продублированной — снимаем, пока снимается.
    for _ in range(3):
        lowered = cleaned.lower()
        for prefix in _LONG_LABELS:
            if lowered.startswith(prefix):
                cleaned = cleaned[len(prefix) :].lstrip(" :\u00a0").strip()
                break
        else:
            for prefix in _SHORT_LABELS:
                if lowered.startswith(prefix + ":"):
                    cleaned = cleaned[len(prefix) + 1 :].strip()
                    break
            else:
                break
    return cleaned.lstrip(" :\u00a0").strip()


def _after_colon(text: str) -> str:
    return text.split(":", 1)[1].strip() if ":" in text else ""


def _parse_title_sheet(rows: list[list[object]]) -> dict[str, str | int | None]:
    """Вытаскивает с титула организацию, центр, год и дату формирования."""
    info: dict[str, str | int | None] = {
        "organization": None,
        "center": None,
        "year": None,
        "report_date": None,
    }
    # Значения лежат в двух строках сразу под блоком «Предоставляет:» — это
    # единственная стабильная опора: часть организаций стирает служебные подписи.
    lines = [next((normalize(c) for c in row if normalize(c)), "") for row in rows]
    anchor = next(
        (i for i, line in enumerate(lines) if line.lower().startswith("предоставляет")),
        None,
    )
    if anchor is not None:
        filled = [line for line in lines[anchor + 1 : anchor + 4] if line]
        if len(filled) >= 1:
            info["organization"] = _strip_label(filled[0]) or None
        if len(filled) >= 2:
            info["center"] = _strip_label(filled[1]) or None

    for line in lines:
        if not line:
            continue
        joined = line
        lowered = joined.lower()

        if info["organization"] is None and any(m in lowered for m in _ORG_MARKERS):
            info["organization"] = _strip_label(joined) or None
        if info["center"] is None and any(m in lowered for m in _CENTER_MARKERS):
            info["center"] = _strip_label(joined) or None
        if info["year"] is None:
            year = re.search(r"\b(20\d{2})\s*год", lowered)
            if year:
                info["year"] = int(year.group(1))
        if info["report_date"] is None and "дата формирования" in lowered:
            date = re.search(r"(\d{2}\.\d{2}\.\d{4})", joined)
            if date:
                info["report_date"] = date.group(1)
    return info


# --- публичный вход -----------------------------------------------------------

@dataclass
class ParsedReport:
    """Результат разбора одной книги."""

    source_file: str
    organization: str | None
    center: str | None
    year: int | None
    report_date: str | None
    values: dict[str, ParsedValue]
    warnings: list[str]

    @property
    def coverage(self) -> float:
        """Доля показателей схемы, для которых найдено фактическое значение."""
        filled = sum(1 for v in self.values.values() if v.fact is not None)
        return filled / len(BY_FORM_CODE) if BY_FORM_CODE else 0.0


def parse_report(path: Path) -> ParsedReport:
    """Разбирает одну книгу отчёта в набор значений показателей."""
    sheets = read_workbook(path)
    title: dict[str, str | int | None] = {}
    values: dict[str, ParsedValue] = {}
    warnings: list[str] = []

    for name, rows in sheets.items():
        role = sheet_role(name)
        if role is None:
            continue
        if role == TITLE:
            title = _parse_title_sheet(rows)
            continue
        parsed, sheet_warnings = _parse_indicator_sheet(rows, role)
        for value in parsed:
            values[value.key] = value
        warnings.extend(sheet_warnings)

    return ParsedReport(
        source_file=path.name,
        organization=title.get("organization"),  # type: ignore[arg-type]
        center=title.get("center"),  # type: ignore[arg-type]
        year=title.get("year"),  # type: ignore[arg-type]
        report_date=title.get("report_date"),  # type: ignore[arg-type]
        values=values,
        warnings=warnings,
    )


def parse_directory(directory: Path) -> list[ParsedReport]:
    """Разбирает все книги в каталоге, пропуская временные файлы Excel (`~$`)."""
    paths = sorted(
        p
        for p in directory.iterdir()
        if p.suffix.lower() in {".xls", ".xlsx"} and not p.name.startswith("~$")
    )
    reports = []
    for path in paths:
        try:
            reports.append(parse_report(path))
        except Exception as exc:  # одна битая книга не должна ронять всю загрузку
            log.error("не удалось разобрать %s: %s", path.name, exc)
    return reports
